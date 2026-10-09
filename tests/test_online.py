from datetime import date
import json
from pathlib import Path
import tempfile
from types import SimpleNamespace
import unittest
from unittest.mock import patch, MagicMock
from openpyxl import load_workbook
import consulta_online as online
import monitor
from exemplos.criar_demo import pdf_ficticio, documento


class OnlineTests(unittest.TestCase):
    def test_windows_usa_edge_visivel_sem_perfil_pessoal(self):
        with patch.object(online.sys,'platform','win32'):
            self.assertEqual(online.opcoes_navegador(None),{'headless':False,'proxy':None,'channel':'msedge'})

    def test_linux_preserva_chromium(self):
        with patch.object(online.sys,'platform','linux'),patch.object(online.shutil,'which',return_value='/usr/bin/chromium'):
            self.assertEqual(online.opcoes_navegador(None),{'headless':True,'proxy':None,'executable_path':'/usr/bin/chromium'})

    def test_403_inicial_interrompe_sem_pesquisar(self):
        browser=MagicMock();page=browser.new_context.return_value.new_page.return_value
        page.goto.return_value.status=403
        with tempfile.TemporaryDirectory() as tmp,patch.object(online,'preparar_filtros') as filtros:
            itens,pendencias,totais=online.coletar(browser,Path(tmp),date(2026,10,9),[],lambda m:None)
        filtros.assert_not_called();self.assertEqual(browser.new_context.call_count,1)
        self.assertEqual(itens,[]);self.assertEqual(totais,{})
        self.assertIn('Nenhum total foi obtido',pendencias[0])

    def test_paginacao_e_erro_nao_sao_zero(self):
        self.assertEqual(online.paginacao('Resultado: 31 Instrumento(s) Coletivo(s) Encontrado(s) - Página 2 de 4'),(31,2,4))
        self.assertEqual(online.paginacao('Resultado: 0 Instrumento(s) Coletivo(s) Encontrado(s)'),(0,1,1))
        for texto in ['', 'HTTP 500', 'Erro de CAPTCHA', 'Resultado: 2 Instrumento(s) Coletivo(s) Encontrado(s) - Página 3 de 2']:
            with self.subTest(texto=texto),self.assertRaises(online.ConsultaIncompleta): online.paginacao(texto)

    def test_registro_nacional_e_solicitacao(self):
        self.assertEqual(monitor.extrair(documento(numero='SRT000001/2026'),date(2026,10,8))['Registro'],'SRT000001/2026')
        self.assertEqual(monitor.extrair(documento(numero='SRT00001/2026'),date(2026,10,8))['Registro'],'SRT00001/2026')
        with self.assertRaises(ValueError): monitor.registro('MR063147/2026')

    def test_http_500_interrompe_sem_zero(self):
        page=MagicMock();page.expect_response.return_value.__enter__.return_value.value.status=500
        with self.assertRaisesRegex(online.ConsultaIncompleta,'HTTP 500'):online.pesquisar(page)
        page.locator.assert_called_once_with('#btnPesquisar')

    def test_mensagem_zero_em_http_500_continua_pendencia(self):
        page=MagicMock();r=page.expect_response.return_value.__enter__.return_value.value
        r.status=500;r.json.return_value={'message':'Nenhum registro encontrado.'}
        with self.assertRaisesRegex(online.ConsultaIncompleta,'Nenhum registro encontrado'):
            online.pesquisar(page)

    def test_coleta_nao_declara_completa_apos_falha_de_um_tipo(self):
        browser=MagicMock()
        context=browser.new_context.return_value;page=context.new_page.return_value
        page.goto.return_value.status=200;page.locator.return_value.count.return_value=0
        with tempfile.TemporaryDirectory() as tmp,patch.object(online,'preparar_filtros'),patch.object(online,'pesquisar',side_effect=[(0,1,1),online.ConsultaIncompleta('Pesquisa oficial respondeu HTTP 500')]):
            items,pendencias,totais=online.coletar(browser,Path(tmp),date(2026,10,8),[],lambda m:None)
        self.assertEqual(items,[]);self.assertIn('HTTP 500',pendencias[0]);self.assertEqual(totais,{'convencao':0})
        self.assertEqual(context.close.call_count,2)

    def test_403_nao_tenta_outro_contexto_para_contornar(self):
        browser=MagicMock();page=browser.new_context.return_value.new_page.return_value
        page.goto.return_value.status=200;page.locator.return_value.count.return_value=0
        with tempfile.TemporaryDirectory() as tmp,patch.object(online,'preparar_filtros'),patch.object(online,'pesquisar',side_effect=online.ConsultaIncompleta('Pesquisa oficial respondeu HTTP 403')):
            _,pendencias,totais=online.coletar(browser,Path(tmp),date(2026,10,8),[],lambda m:None)
        self.assertEqual(browser.new_context.call_count,1)
        self.assertIn('HTTP 403',pendencias[0]);self.assertEqual(totais,{})

    def test_execucao_sem_manifesto_usuario_preserva_e_deduplica(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);wb,ws,_,_=monitor.abrir_controle(None)
            ws.append(['AL000099/2026','01/01/2026','Anterior','AL','antes.pdf','fonte','=1+1'])
            controle=p/'original.xlsx';wb.save(controle);original=controle.read_bytes()
            def coleta(browser,docs,hoje,eventos,progresso,inicio=None,catalogo=None):
                pdf_ficticio(docs/'novo.pdf',documento())
                return [{'registro':'AL000001/2026','arquivo':'documentos/novo.pdf','fonte':online.URL}],[],{'convencao':1,'termoAditivoConvecao':0}
            with patch.object(online,'coletar',side_effect=coleta),patch.object(online,'hoje_alagoas',return_value=date(2026,10,8)),patch('playwright.sync_api.sync_playwright'):
                r=online.executar_online(controle,p/'saidas')
                self.assertTrue(r['completa']);self.assertEqual(r['codigo'],0)
                self.assertEqual(controle.read_bytes(),original)
                planilha=load_workbook(r['controle']);self.assertEqual(planilha.active.max_row,3)
                self.assertEqual(planilha.active['G2'].value,'=1+1')
                self.assertIn('Modo online',r['resumo'])
                r2=online.executar_online(r['controle'],p/'saidas')
                self.assertTrue(r2['completa']);self.assertEqual(load_workbook(r2['controle']).active.max_row,3)
                self.assertIn('Novos instrumentos: 0',r2['resumo'])

    def test_falha_previa_gera_relatorios_sem_envio_de_controle(self):
        with tempfile.TemporaryDirectory() as tmp,patch('playwright.sync_api.sync_playwright') as browser:
            p=Path(tmp);r=online.executar_online(p/'ausente.xlsx',p/'saidas')
            self.assertFalse(r['completa']);self.assertEqual(r['codigo'],2)
            browser.assert_not_called()
            self.assertTrue((r['pasta']/'log_convencoes_AL.txt').is_file())


if __name__ == '__main__':unittest.main()
