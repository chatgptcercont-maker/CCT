from datetime import date
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch, MagicMock
from openpyxl import load_workbook
import monitor
import consulta_online as online
from exemplos.criar_demo import pdf_ficticio, documento


class AnualTests(unittest.TestCase):
    hoje=date(2026,10,9)

    def test_linha_com_referencia_usa_registro_do_documento(self):
        page=MagicMock();row=MagicMock();doc=MagicMock()
        row.inner_text.return_value='AL000001/2026 referência AL000099/2025'
        row.get_attribute.return_value='MR000001/2026'
        link=MagicMock();link.count.return_value=1
        link.get_attribute.return_value=online.URL
        download=MagicMock();download.count.return_value=0
        row.locator.return_value.filter.side_effect=[link,download]
        page.expect_popup.return_value.__enter__.return_value.value=doc
        doc.locator.return_value.inner_text.return_value=documento(data='01/01/2026')
        with tempfile.TemporaryDirectory() as tmp:
            def imprimir(**args):pdf_ficticio(Path(args['path']),documento(data='01/01/2026'))
            doc.pdf.side_effect=imprimir
            item=online.salvar_instrumento(page,row,Path(tmp),self.hoje,[],date(2026,1,1))
        self.assertEqual(item['registro'],'AL000001/2026')
        doc.close.assert_called_once()

    def test_periodo_2026_nao_inclui_datas_futuras(self):
        self.assertEqual(online.periodo_consulta(self.hoje,2026),(date(2026,1,1),self.hoje))
        self.assertEqual(online.periodo_consulta(self.hoje,2025),(date(2025,1,1),date(2025,12,31)))
        with self.assertRaises(ValueError):online.periodo_consulta(self.hoje,2027)

    def test_janeiro_aceito_no_anual_e_excluido_no_diario(self):
        texto=documento(data='01/01/2026')
        self.assertEqual(monitor.extrair(texto,self.hoje,date(2026,1,1))['Data de registro'],'01/01/2026')
        with self.assertRaisesRegex(ValueError,'fora do período'):monitor.extrair(texto,self.hoje)

    def test_limites_tipo_e_ano_na_data_de_registro(self):
        for d in ['31/12/2025','10/10/2026']:
            with self.assertRaises(ValueError):monitor.extrair(documento(data=d),self.hoje,date(2026,1,1))
        for tipo in ['ACORDO COLETIVO DE TRABALHO','TERMO ADITIVO DE ACORDO COLETIVO DE TRABALHO']:
            with self.assertRaises(ValueError):monitor.extrair(documento(data='01/01/2026',tipo=tipo),self.hoje,date(2026,1,1))
        dados=monitor.extrair(documento(data='09/10/2026',tipo='TERMO ADITIVO DE CONVENCAO COLETIVA DE TRABALHO'),self.hoje,date(2026,1,1))
        self.assertIn('Aditivo',dados['Tipo'])

    def test_relatorio_anual_inclui_registro_ja_existente(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);wb,ws,_,_=monitor.abrir_controle(None)
            ws.append(['AL000001/2026','01/01/2026','Antigo','AL','origem.pdf','fonte','hash']);wb.save(p/'original.xlsx')
            antes=(p/'original.xlsx').read_bytes()
            def coletar(browser,docs,hoje,eventos,progresso,inicio=None,catalogo=None):
                self.assertEqual(inicio,date(2026,1,1))
                pdf_ficticio(docs/'janeiro.pdf',documento(data='01/01/2026'))
                return [{'registro':'AL000001/2026','arquivo':'documentos/janeiro.pdf','fonte':online.URL}],[],{tipo+':'+area: (1 if tipo=='convencao' else 0) for tipo,area in online.PESQUISAS}
            with patch.object(online,'hoje_alagoas',return_value=self.hoje),patch.object(online,'coletar',side_effect=coletar),patch('playwright.sync_api.sync_playwright'):
                r=online.executar_online(p/'original.xlsx',p/'saidas',ano=2026)
            self.assertTrue(r['completa']);self.assertEqual(antes,(p/'original.xlsx').read_bytes())
            self.assertEqual(load_workbook(r['controle']).active.max_row,2)
            rel=load_workbook(r['pasta']/'instrumentos_consultados_AL.xlsx')
            self.assertEqual(rel['Instrumentos'].max_row,2)
            self.assertEqual(rel['Instrumentos']['A2'].value,'AL000001/2026')
            self.assertIn('Novos instrumentos: 0',r['resumo'])
            self.assertIn('AL000001/2026',(r['pasta']/'instrumentos_consultados_AL.txt').read_text())

    def test_pendencia_documento_nao_preenche_dados_sem_original(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);out=p/'resultado';out.mkdir()
            catalogo=[{'Registro':'AL000002/2026','Solicitação':'MR000002/2026','Pesquisa':'convencao','Status':'Documento indisponível'}]
            online.gerar_relatorio(catalogo,[],p,out,date(2026,1,1),self.hoje,False)
            wb=load_workbook(out/'instrumentos_consultados_AL.xlsx')
            self.assertEqual(wb['Instrumentos'].max_row,1)
            self.assertEqual(wb['Pendências']['A2'].value,'AL000002/2026')
            self.assertIsNone(wb['Pendências']['B2'].value)
            self.assertIn('CONSULTA COM PENDÊNCIAS',(out/'instrumentos_consultados_AL.txt').read_text())

    def test_falha_de_acesso_aparece_no_relatorio_sem_documentos(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);out=p/'resultado';out.mkdir()
            online.gerar_relatorio([],[],p,out,date(2026,1,1),self.hoje,False,['HTTP 403: acesso recusado'])
            self.assertIn('FALHA DE CONSULTA: HTTP 403',(out/'instrumentos_consultados_AL.txt').read_text())
            wb=load_workbook(out/'instrumentos_consultados_AL.xlsx')
            self.assertEqual(wb['Consulta'].cell(wb['Consulta'].max_row,2).value,'HTTP 403: acesso recusado')


if __name__ == '__main__':unittest.main()
