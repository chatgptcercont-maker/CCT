import hashlib
import json
from pathlib import Path
import tempfile
import unittest
from datetime import date
from unittest.mock import patch
from urllib.error import URLError
from openpyxl import load_workbook
import monitor
from exemplos.criar_demo import documento, pdf_ficticio


class MonitorTests(unittest.TestCase):
    hoje=date(2026,10,8)

    def test_tipos(self):
        self.assertEqual(monitor.extrair(documento(),self.hoje)['Registro'],'AL000001/2026')
        self.assertIn('Aditivo',monitor.extrair(documento(tipo='TERMO ADITIVO DE CONVENCAO COLETIVA DE TRABALHO'),self.hoje)['Tipo'])
        for tipo in ['ACORDO COLETIVO DE TRABALHO','TERMO ADITIVO DE ACORDO COLETIVO DE TRABALHO','TERMO ADITIVO']:
            with self.subTest(tipo=tipo), self.assertRaises(ValueError): monitor.extrair(documento(tipo=tipo),self.hoje)

    def test_periodo(self):
        for data in ['02/10/2026','08/10/2026']: monitor.extrair(documento(data=data),self.hoje)
        for data in ['01/10/2026','09/10/2026','99/10/2026']:
            with self.subTest(data=data), self.assertRaises(ValueError): monitor.extrair(documento(data=data),self.hoje)

    def test_territorio(self):
        for texto in [documento(local='Recife/PE'),documento().replace('ABRANGENCIA TERRITORIAL','AUSENTE')]:
            with self.assertRaises(ValueError): monitor.extrair(texto,self.hoje)

    def test_campos(self):
        for texto in [documento().replace('DATA DE REGISTRO NO MTE:','DATA:'), documento()+'\nNUMERO DE REGISTRO NO MTE: AL000002/2026']:
            with self.assertRaises(ValueError): monitor.extrair(texto,self.hoje)

    def test_pdf_planilha_preservacao_repeticao(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp); pdf_ficticio(p/'a.pdf',documento())
            wb,ws,_,_=monitor.abrir_controle(None)
            ws.append(['AL000099/2026','01/01/2026','Antigo','AL','antigo.pdf','fonte','=1+1']); wb.save(p/'original.xlsx')
            original=(p/'original.xlsx').read_bytes()
            item={'arquivo':'a.pdf','fonte':monitor.CONSULTA,'registro':'AL000001/2026'}
            (p/'dados.json').write_text(json.dumps({'consulta_completa':True,'instrumentos':[item,item]}))
            self.assertEqual(monitor.executar(p/'dados.json',p/'original.xlsx',p/'out',self.hoje),0)
            self.assertEqual(original,(p/'original.xlsx').read_bytes())
            ws=load_workbook(p/'out'/'controle_convencoes_AL.xlsx').active
            self.assertEqual(ws.max_row,3); self.assertEqual(ws['G2'].value,'=1+1')
            self.assertEqual(ws['A3'].value,'AL000001/2026')
            self.assertEqual(ws['G3'].value,hashlib.sha256((p/'a.pdf').read_bytes()).hexdigest())
            self.assertEqual(monitor.executar(p/'dados.json',p/'out'/'controle_convencoes_AL.xlsx',p/'out2',self.hoje),0)
            self.assertEqual(load_workbook(p/'out2'/'controle_convencoes_AL.xlsx').active.max_row,3)
            with self.assertRaises(FileExistsError): monitor.executar(p/'dados.json',None,p/'out',self.hoje)

    def test_pendencias(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp); (p/'dados.json').write_text(json.dumps({'consulta_completa':False,'instrumentos':[
                {'status':'captcha'},{'status':'documento_indisponivel'},
                {'arquivo':'ausente.pdf','fonte':monitor.CONSULTA},{'arquivo':'../fora.pdf','fonte':monitor.CONSULTA}]}))
            self.assertEqual(monitor.executar(p/'dados.json',None,p/'out',self.hoje),2)
            self.assertEqual(load_workbook(p/'out'/'controle_convencoes_AL.xlsx').active.max_row,1)
            self.assertIn('captcha',(p/'out'/'log_convencoes_AL.txt').read_text())
            with patch('urllib.request.urlopen',side_effect=URLError('403 Forbidden')):
                self.assertEqual(monitor.consultar(p/'consulta'),2)
            self.assertIn('403 Forbidden',(p/'consulta'/'log_convencoes_AL.txt').read_text())

    def test_incompatibilidade_e_pdf_vazio(self):
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp); pdf_ficticio(p/'vazio.pdf','')
            with self.assertRaises(ValueError): monitor.ler_pdf(p/'vazio.pdf')
            wb,ws,_,_=monitor.abrir_controle(None); ws['A1']='Outra coluna'; wb.save(p/'controle.xlsx')
            with self.assertRaises(ValueError): monitor.abrir_controle(p/'controle.xlsx')

    def test_fonte(self):
        for url in ['https://www3.mte.gov.br.evil.test/','http://www3.mte.gov.br/','https://senha@www3.mte.gov.br/']:
            self.assertFalse(monitor.fonte_oficial(url))


if __name__ == '__main__': unittest.main()
