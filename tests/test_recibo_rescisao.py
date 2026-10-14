from decimal import Decimal
from pathlib import Path
import tempfile
import unittest
from pypdf import PdfReader
from recibo_rescisao import preparar_recibo,valor_monetario,exportar_recibo,texto_recibo


def dados():
    return {'empresa':'EMPRESA FICTICIA','trabalhador':'TRABALHADOR FICTICIO','local':'Maceió',
            'data_recibo':'14/10/2026','admissao':'01/01/2026','desligamento':'10/10/2026'}


class ReciboTests(unittest.TestCase):
    def test_valores_brasileiros_e_decimais_sem_arredondar_silenciosamente(self):
        self.assertEqual(valor_monetario('R$ 1.250,50'),Decimal('1250.50'))
        self.assertEqual(valor_monetario('1.000'),Decimal('1000.00'))
        self.assertEqual(valor_monetario('1250.50'),Decimal('1250.50'))
        for v in ['-10','10,001','abc','1e3','']:
            with self.assertRaises(ValueError):valor_monetario(v)

    def test_soma_apenas_as_verbas_informadas(self):
        r=preparar_recibo(dados(),[{'tipo':'Provento','descricao':'Verba manual','valor':'0,10'},
                                  {'tipo':'Provento','descricao':'Outra verba manual','valor':'0,20'},
                                  {'tipo':'Desconto','descricao':'Desconto informado','valor':'0,05'}])
        self.assertEqual(r['liquido'],Decimal('0.25'))
        self.assertEqual(len(r['itens']),3)
        self.assertEqual([i['descricao'] for i in r['itens']],['Verba manual','Outra verba manual','Desconto informado'])

    def test_dados_incompletos_datas_e_total_negativo_impedem_emissao(self):
        itens=[{'tipo':'Provento','descricao':'Verba manual','valor':'100'}]
        d=dados();d['empresa']=''
        with self.assertRaises(ValueError):preparar_recibo(d,itens)
        d=dados();d['admissao']='20/10/2026'
        with self.assertRaises(ValueError):preparar_recibo(d,itens)
        with self.assertRaises(ValueError):preparar_recibo(dados(),[])
        with self.assertRaises(ValueError):preparar_recibo(dados(),itens+[{'tipo':'Desconto','descricao':'Desconto','valor':'101'}])

    def test_pdf_reproduz_valores_e_nao_sobrescreve_existentes(self):
        r=preparar_recibo(dados(),[{'tipo':'Provento','descricao':'Saldo informado','valor':'1.250,50'},
                                  {'tipo':'Desconto','descricao':'Adiantamento informado','valor':'50,00'}])
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'recibo.pdf';exportar_recibo(r,p)
            conteudo='\n'.join(x.extract_text() or '' for x in PdfReader(p).pages)
            self.assertIn('R$ 1.200,50',conteudo);self.assertIn('EMPRESA FICTICIA',conteudo)
            self.assertIn('TERMO DE RESCISÃO',conteudo)
            original=p.read_bytes()
            with self.assertRaises(ValueError):exportar_recibo(r,p)
            self.assertEqual(original,p.read_bytes())
            with self.assertRaises(ValueError):exportar_recibo(r,Path(tmp)/'controle.xlsx')

    def test_modelo_duas_paginas_e_dados_consistentes(self):
        from modelo_rescisao import exportar_excel
        from openpyxl import load_workbook
        d=dados();d.update({'emitir_quitacao':True,'data_pagamento':'14/10/2026','causa':'Causa informada manualmente'})
        r=preparar_recibo(d,[{'tipo':'Provento','descricao':'Verba manual','valor':'1000'}])
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp);exportar_recibo(r,p/'modelo.pdf')
            leitor=PdfReader(p/'modelo.pdf');self.assertEqual(len(leitor.pages),2)
            self.assertIn('EMPRESA FICTICIA',leitor.pages[0].extract_text())
            self.assertIn('EMPRESA FICTICIA',leitor.pages[1].extract_text())
            self.assertIn('R$ 1.000,00',leitor.pages[1].extract_text())
            exportar_excel(r,p/'modelo.xlsx');w=load_workbook(p/'modelo.xlsx')
            self.assertEqual(w.worksheets[0]['C6'].value,w.worksheets[1]['C5'].value)
            self.assertEqual(w.worksheets[0]['A24'].value,w.worksheets[1]['A13'].value)
            self.assertEqual(w.worksheets[0]['F49'].value,'R$ 1.000,00')

    def test_quitacao_exige_data_pagamento_e_nao_e_gerada_por_padrao(self):
        from modelo_rescisao import folhas
        itens=[{'tipo':'Provento','descricao':'Verba manual','valor':'100'}]
        r=preparar_recibo(dados(),itens);self.assertEqual(len(folhas(r)),1)
        d=dados();d['emitir_quitacao']=True
        with self.assertRaises(ValueError):preparar_recibo(d,itens)

    def test_excel_preserva_texto_e_observacoes_sem_formulas(self):
        from modelo_rescisao import exportar_excel
        from openpyxl import load_workbook
        d=dados();d.update(empresa='=1+1',observacoes='Observação fictícia')
        r=preparar_recibo(d,[{'tipo':'Provento','descricao':'Verba manual','valor':'100'}])
        with tempfile.TemporaryDirectory() as tmp:
            p=Path(tmp)/'modelo.xlsx';exportar_excel(r,p)
            w=load_workbook(p);self.assertEqual(w.active['C6'].data_type,'s')
            self.assertEqual(w.active['C6'].value,'=1+1')
            self.assertIn('Observação fictícia',w.active['A31'].value)
            original=p.read_bytes()
            with self.assertRaises(ValueError):exportar_excel(r,p)
            self.assertEqual(p.read_bytes(),original)

    def test_modelo_sem_identificadores_pessoais_e_limite_de_verbas(self):
        import json,re
        from modelo_rescisao import MODELO,folhas
        spec=json.loads(MODELO.read_text(encoding='utf-8'))
        for aba in spec['abas']:
            for rotulo in aba['rotulos'].values():
                self.assertIsNone(re.search(r'\d{3}\.\d{3}\.\d{3}-\d{2}|\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}',rotulo))
        itens=[{'tipo':'Provento','descricao':'Verba fictícia','valor':'100'}]*19
        with self.assertRaises(ValueError):folhas(preparar_recibo(dados(),itens))


if __name__=='__main__':unittest.main()
