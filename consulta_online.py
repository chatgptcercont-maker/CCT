"""Coleta pela navegação pública normal do Mediador, sem contornar proteções."""
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from urllib.parse import urljoin, urlsplit
import uuid
from zoneinfo import ZoneInfo

from diagnosticar_mediador import URL, proxy_do_ambiente, classificar_erro
import monitor

TIPOS = ('convencao', 'termoAditivoConvecao')
PESQUISAS = tuple((tipo, area) for tipo in TIPOS for area in ('estado', 'municipios'))


class ConsultaIncompleta(RuntimeError):
    pass


def opcoes_navegador(proxy):
    """No Windows, usa o Edge instalado com janela e sessão separada."""
    options = {'headless': sys.platform != 'win32', 'proxy': proxy}
    if sys.platform == 'win32':
        options['channel'] = 'msedge'
    else:
        executavel = shutil.which('chromium') or shutil.which('chromium-browser')
        if executavel:
            options['executable_path'] = executavel
    return options


def captcha_interativo(page):
    desafios = page.locator('iframe[src*="bframe"], iframe[src*="hcaptcha"]')
    return any(desafios.nth(i).is_visible() for i in range(desafios.count()))


def hoje_alagoas():
    return datetime.now(ZoneInfo('America/Maceio')).date()


def periodo_consulta(hoje, ano=None, ontem=False):
    if ontem:
        if ano is not None:
            raise ValueError('Escolha o ano ou ontem, não os dois períodos')
        dia = hoje-timedelta(days=1)
        return dia, dia
    if ano is None:
        return hoje-timedelta(days=6), hoje
    if not 1900 <= ano <= hoje.year:
        raise ValueError('Informe um ano já iniciado, entre 1900 e o ano atual')
    return date(ano, 1, 1), min(hoje, date(ano, 12, 31))


def paginacao(texto):
    n = monitor.normalizar(texto)
    m = re.search(r'RESULTADO:\s*(\d+)\s+INSTRUMENTO\(S\).*?PAGINA\s+(\d+)\s+DE\s+(\d+)', n, re.S)
    if m:
        total, atual, paginas = map(int, m.groups())
        if total < 0 or atual < 1 or paginas < atual or (total and not paginas):
            raise ConsultaIncompleta('Contagem de resultados inconsistente')
        return total, atual, paginas
    # Somente mensagem explicitamente observável; página vazia nunca significa zero.
    if re.fullmatch(r'\s*RESULTADO:\s*(?:NENHUM INSTRUMENTO COLETIVO ENCONTRADO|0 INSTRUMENTO\(S\) COLETIVO\(S\) ENCONTRADO\(S\))\s*', n):
        return 0, 1, 1
    raise ConsultaIncompleta('Formato do resultado não reconhecido; não é possível confirmar a paginação')


def preparar_filtros(page, hoje, tipo, inicio=None, area='estado'):
    page.locator('#cboTPRequerimento').select_option(tipo)
    page.locator('#chkPeriodoRegistro').check()
    page.locator('#txtDTInicioRegistro').fill((inicio or hoje-timedelta(days=6)).strftime('%d/%m/%Y'))
    page.locator('#txtDTFimRegistro').fill(hoje.strftime('%d/%m/%Y'))
    page.locator('#cboSTVigencia').select_option('2')
    page.locator('#cboUFRegistro').select_option('')
    for campo in ['chkMunicipal', 'chkIntermunicipal', 'chkEstadual', 'chkInterestadual', 'chkNacional']:
        checkbox = page.locator('#'+campo)
        if not checkbox.is_checked():
            # Cada alteração dispara uma limpeza assíncrona da seleção territorial.
            # Aguarda essa limpeza antes de adicionar UF ou municípios.
            with page.expect_response(lambda r: '/getPartAbrangenciaRemoverTodos' in r.url, timeout=30000) as limpeza:
                checkbox.check()
            if limpeza.value.status != 200:
                raise ConsultaIncompleta(f'Limpeza do filtro territorial respondeu HTTP {limpeza.value.status}')
    page.locator('#cboUFAbrangenciaTerritorial').select_option('AL')
    if area == 'municipios':
        selecionar_municipios_al(page)
        return
    if area != 'estado':
        raise ValueError('Área de pesquisa desconhecida')
    with page.expect_response(lambda r: '/getPartAbrangenciaAddUF' in r.url, timeout=30000) as espera:
        page.locator('#btnAdicionarUF').click()
    if espera.value.status != 200:
        raise ConsultaIncompleta(f'Seleção territorial falhou: HTTP {espera.value.status}')


def selecionar_municipios_al(page):
    """Usa Todos disponíveis, não apenas os municípios da inicial exibida."""
    with page.expect_response(lambda r: '/getPartMunicipiosPelaUF' in r.url, timeout=30000) as municipios:
        page.locator('#btnSelecionarMunicipio').click()
    if municipios.value.status != 200:
        raise ConsultaIncompleta(f'Seletor de municípios respondeu HTTP {municipios.value.status}')
    dialogo = page.locator('#divPopUpAbrangenciaMunicipios')
    dialogo.wait_for(state='visible', timeout=15000)
    if not re.search(r'\bUF:\s*AL\b', dialogo.inner_text()):
        raise ConsultaIncompleta('O seletor de municípios não confirmou a UF AL')
    # O clique substitui o HTML do checkbox; check() pode falhar ao conferir
    # o elemento antigo. Aguarda as respostas nativas de seleção e gravação.
    with page.expect_response(lambda r: '/getPartAbrangenciaGravarBufferDaUFAtiva' in r.url, timeout=30000) as gravacao:
        with page.expect_response(lambda r: '/getPartMunicipiosUFMarcada' in r.url, timeout=30000) as selecao:
            page.locator('#chkTodosDisponiveisAbrangencia').click()
        if selecao.value.status != 200:
            raise ConsultaIncompleta(f'Seleção de todos os municípios respondeu HTTP {selecao.value.status}')
    if gravacao.value.status != 200:
        raise ConsultaIncompleta(f'Gravação dos municípios respondeu HTTP {gravacao.value.status}')
    page.locator('.ui-dialog').filter(has=dialogo).get_by_text('Fechar', exact=True).click()
    dialogo.wait_for(state='hidden', timeout=15000)


def pesquisar(page):
    # O código original do botão obtém o reCAPTCHA. Não interceptamos tokens.
    with page.expect_response(lambda r: '/ConsultarInstColetivo/getConsultaAvancada' in r.url, timeout=45000) as espera:
        page.locator('#btnPesquisar').click()
    resposta = espera.value
    if resposta.status != 200:
        detalhe = ''
        # O portal usa HTTP 500 também para notificar ausência de registros.
        # Preservamos o erro e a mensagem literal, sem converter falhas em sucesso.
        if resposta.status == 500:
            try:
                if resposta.json().get('message') == 'Nenhum registro encontrado.':
                    detalhe = ': portal informou "Nenhum registro encontrado."; resposta HTTP de erro preservada'
            except (ValueError, AttributeError, TypeError):
                pass
        raise ConsultaIncompleta(f'Pesquisa oficial respondeu HTTP {resposta.status}' + detalhe)
    page.locator('#divExibirConsultaDetalhada').wait_for(state='visible', timeout=15000)
    return paginacao(page.locator('#divConsultaDetalhada h2').inner_text())


def salvar_instrumento(page, row, pasta, hoje, eventos, inicio=None):
    regs = re.findall(r'\b' + monitor.REGEX_REGISTRO + r'\b', row.inner_text())
    # A linha pode mencionar instrumentos anteriores. O cabeçalho do documento
    # oficial é a referência para identificar o registro consultado.
    reg = monitor.registro(regs[0]) if len(set(regs)) == 1 else None
    identidade = row.get_attribute('indice') or ''
    if not re.fullmatch(r'MR\d{6}/\d{4}', identidade):
        raise ConsultaIncompleta('Solicitação oficial não reconhecida')
    stem = (reg or identidade).replace('/', '_')
    link = row.locator('a').filter(has_text='Visualizar Instrumento Coletivo')
    if link.count() != 1:
        raise ConsultaIncompleta(f'{reg}: link de visualização não reconhecido')
    fonte = urljoin(URL, link.get_attribute('href') or '')
    if not monitor.fonte_oficial(fonte):
        raise ConsultaIncompleta(f'{reg}: visualização fora da fonte oficial')
    original = None
    botao_download = row.locator('a').filter(has_text='Download')
    if botao_download.count() == 1:
        paginas_anteriores = set(page.context.pages)
        try:
            with page.expect_download(timeout=30000) as espera:
                botao_download.click()
            download = espera.value
            if urlsplit(download.url).hostname != urlsplit(URL).hostname:
                raise ConsultaIncompleta('Download redirecionado para origem não confirmada')
            ext = Path(download.suggested_filename).suffix.lower()
            if ext not in {'.pdf', '.doc', '.docx'}:
                raise ConsultaIncompleta('Formato de download oficial não reconhecido')
            original = pasta / (stem+'_original'+ext)
            download.save_as(original)
            if original.stat().st_size > 30*1024*1024:
                raise ConsultaIncompleta('Documento original excede 30 MB')
            eventos.append(f'{reg}: original oficial salvo em {original.name}; SHA256 '+hashlib.sha256(original.read_bytes()).hexdigest())
        except Exception as exc:
            eventos.append(f'{reg}: download indisponível ({type(exc).__name__}); tentando visualização oficial')
            original = None
        finally:
            for auxiliar in page.context.pages:
                if auxiliar not in paginas_anteriores and not auxiliar.is_closed():
                    auxiliar.close()
    # A visualização gera token de segurança e POST por meio do código do portal.
    with page.expect_popup(timeout=30000) as popup:
        link.click()
    doc = popup.value
    try:
        doc.wait_for_function("document.body && document.body.innerText.toUpperCase().includes('NÚMERO DE REGISTRO NO MTE')", timeout=30000)
        original_texto = doc.locator('body').inner_text()
        dados = monitor.extrair(original_texto, hoje, inicio)
        if reg is not None and dados['Registro'] != reg:
            raise ConsultaIncompleta('Registro da visualização diverge do resultado de pesquisa')
        reg = dados['Registro']
        if original is not None and original.suffix == '.pdf':
            pdf = original
        else:
            pdf = pasta / (stem+'_impressao.pdf')
            doc.pdf(path=str(pdf), format='A4', print_background=True)
            eventos.append(f'{reg}: PDF de impressão gerado da visualização oficial; não é PDF original oferecido pelo portal')
        campos_pdf = monitor.extrair(monitor.ler_pdf(pdf), hoje, inicio)
        if any(campos_pdf[c] != dados[c] for c in ['Registro', 'Data de registro', 'Tipo']):
            raise ConsultaIncompleta('Campos do PDF divergem do documento visualizado')
        return {'registro': reg, 'arquivo': 'documentos/'+pdf.name, 'fonte': fonte, 'campos': campos_pdf}
    finally:
        doc.close()


def coletar(browser, pasta, hoje, eventos, progresso, inicio=None, catalogo=None):
    instrumentos, pendencias, totais = [], [], {}
    catalogo = catalogo if catalogo is not None else []
    catalogados = set()
    for tipo, area in PESQUISAS:
        pesquisa = tipo + ':' + area
        progresso({'total': None, 'examinados': 0, 'pesquisa': pesquisa})
        progresso('Pesquisando '+('convenções' if tipo == 'convencao' else 'aditivos de convenções')+
                  (' — estado de Alagoas' if area == 'estado' else ' — todos os municípios de Alagoas')+'…')
        # Sessões independentes impedem que a abrangência/session de uma pesquisa contamine a outra.
        context = browser.new_context(accept_downloads=True)
        context.on('response', lambda r: eventos.append(f'HTTP {r.status}: {urlsplit(r.url).path}')
                   if r.status >= 400 and urlsplit(r.url).hostname == urlsplit(URL).hostname else None)
        page = context.new_page()
        try:
            resposta = page.goto(URL, wait_until='domcontentloaded', timeout=45000)
            if resposta is None or resposta.status != 200:
                if resposta is not None and resposta.status == 403:
                    raise ConsultaIncompleta('HTTP 403: acesso recusado antes da pesquisa. Nenhum total foi obtido; '
                                             'não significa ausência de convenções. A origem da recusa não foi confirmada. '
                                             'Abra o endereço oficial no navegador para comparar o acesso. Não houve tentativa de contornar o bloqueio')
                raise ConsultaIncompleta('Página de consulta não disponível: HTTP '+str(resposta.status if resposta else 'sem resposta'))
            if page.locator('input[type="password"]').count():
                raise ConsultaIncompleta('Autenticação exigida; acesso automático interrompido')
            preparar_filtros(page, hoje, tipo, inicio, area=area)
            total, atual, paginas = pesquisar(page)
            eventos.append(f'{pesquisa}: resultado oficial {total} instrumento(s), {paginas} página(s)')
            vistos = set()
            progresso({'total': total, 'examinados': 0, 'pesquisa': pesquisa})
            if paginas > 200:
                raise ConsultaIncompleta('Mais de 200 páginas; consulta interrompida sem presumir completude')
            for numero in range(1, paginas+1):
                if numero != atual:
                    raise ConsultaIncompleta('Página recebida fora da sequência')
                rows = page.locator('#grdInstrumentos tr[indice]')
                for i in range(rows.count()):
                    row = rows.nth(i)
                    identidade = row.get_attribute('indice')
                    if not identidade or identidade in vistos:
                        raise ConsultaIncompleta('Identificador de instrumento ausente/repetido na paginação')
                    vistos.add(identidade)
                    if identidade in catalogados:
                        eventos.append(f'{pesquisa}: {identidade} já examinado em outra pesquisa; sem duplicar')
                        progresso({'total': total, 'examinados': len(vistos), 'pesquisa': pesquisa})
                        continue
                    catalogados.add(identidade)
                    entrada = {'Solicitação': identidade, 'Pesquisa': pesquisa, 'Status': 'Documento não confirmado'}
                    regs = re.findall(r'\b' + monitor.REGEX_REGISTRO + r'\b', row.inner_text())
                    if len(set(regs)) == 1:
                        entrada['Registro'] = regs[0]
                    catalogo.append(entrada)
                    try:
                        item = salvar_instrumento(page, row, pasta, hoje, eventos, inicio)
                        instrumentos.append(item)
                        entrada.update(item['campos'])
                        entrada.update({'Documento': item['arquivo'], 'Fonte': item['fonte'], 'Status': 'Documento lido e validado'})
                        progresso('Documento lido: '+item['registro'])
                    except Exception as exc:
                        motivo = str(exc) if isinstance(exc, (ConsultaIncompleta, ValueError)) else type(exc).__name__
                        eventos.append(f'{identidade}: documento indisponível ou rejeitado: {motivo}')
                        entrada['Status'] = motivo
                        # Exclusões documentadas não são uma falha na coleta.
                        if not motivo.startswith('Excluído:'):
                            pendencias.append(tipo+': '+identidade+': '+motivo)
                        progresso(f'{identidade}: {motivo}')
                    progresso({'total': total, 'examinados': len(vistos), 'pesquisa': pesquisa})
                if numero < paginas:
                    proxima = page.locator('#grdInstrumentos > table > tfoot a').filter(has_text=re.compile(r'^\s*'+str(numero+1)+r'\s*$'))
                    if proxima.count() != 1:
                        raise ConsultaIncompleta('Controle da próxima página não reconhecido')
                    with page.expect_response(lambda r: '/getConsultaAvancada' in r.url, timeout=45000) as espera:
                        proxima.click()
                    if espera.value.status != 200:
                        raise ConsultaIncompleta(f'Paginação respondeu HTTP {espera.value.status}')
                    page.wait_for_function("n => document.querySelector('#divConsultaDetalhada h2').innerText.includes('Página '+n+' de')", arg=numero+1, timeout=15000)
                    novo_total, atual, novo_paginas = paginacao(page.locator('#divConsultaDetalhada h2').inner_text())
                    if (novo_total, novo_paginas) != (total, paginas):
                        raise ConsultaIncompleta('Resultado mudou durante paginação; consultar novamente')
            if len(vistos) != total:
                raise ConsultaIncompleta(f'Total oficial {total} difere dos {len(vistos)} registros examinados')
            totais[pesquisa] = total
        except Exception as exc:
            motivo = str(exc) if isinstance(exc, ConsultaIncompleta) else classificar_erro(str(exc))+': '+type(exc).__name__
            if not page.is_closed() and captcha_interativo(page):
                motivo = 'CAPTCHA interativo detectado; consulta interrompida sem resolver ou contornar o desafio'
            pendencias.append(pesquisa+': '+motivo)
            eventos.append(pesquisa+': '+motivo+'; sem contornar a proteção ou repetir a pesquisa')
            progresso(pesquisa+': '+motivo)
            if any(s in motivo for s in ['HTTP 401', 'HTTP 403', 'CAPTCHA interativo', 'Autenticação exigida', 'BLOQUEIO_DO_TUNEL_PROXY', 'ERRO_TLS']):
                break
        finally:
            context.close()
    return instrumentos, pendencias, totais


def gerar_relatorio(catalogo, instrumentos, pasta, resultado, inicio, fim, completo, falhas_consulta=None):
    """Lista todos os documentos da consulta, inclusive os já presentes no controle."""
    from openpyxl import Workbook
    import csv
    # Inclui resultados de coletores compatíveis que não preencheram o catálogo.
    identificados = {e.get('Registro') for e in catalogo}
    for item in instrumentos:
        if item['registro'] not in identificados:
            dados = monitor.extrair(monitor.ler_pdf(pasta/item['arquivo']), fim, inicio)
            catalogo.append({**dados, 'Documento': item['arquivo'], 'Fonte': item['fonte'], 'Status': 'Documento lido e validado'})
            identificados.add(item['registro'])
    colunas = monitor.COLUNAS + ['Status', 'Pesquisa', 'Solicitação']
    validos, pendentes = [], []
    for entrada in catalogo:
        if entrada.get('Status', '').startswith('Excluído:'):
            continue
        if entrada.get('Documento'):
            entrada['SHA256'] = hashlib.sha256((pasta/entrada['Documento']).read_bytes()).hexdigest()
        if entrada.get('Status') == 'Documento lido e validado':
            validos.append(entrada)
        else:
            pendentes.append(entrada)
    validos.sort(key=lambda d: (datetime.strptime(d['Data de registro'], '%d/%m/%Y'), d['Registro']))
    wb = Workbook(); ws = wb.active; ws.title = 'Instrumentos'
    pend = wb.create_sheet('Pendências')
    for aba, linhas in [(ws, validos), (pend, pendentes)]:
        aba.append(colunas)
        for item in linhas:
            aba.append([("'"+str(item[c]) if str(item.get(c, '')).startswith(('=', '+', '-', '@')) else item.get(c)) for c in colunas])
        aba.freeze_panes = 'A2'; aba.auto_filter.ref = aba.dimensions
    estado = wb.create_sheet('Consulta')
    estado.append(['Informação', 'Valor'])
    estado.append(['Data inicial de registro', inicio.strftime('%d/%m/%Y')])
    estado.append(['Data final de registro', fim.strftime('%d/%m/%Y')])
    estado.append(['Resultado', 'Consulta completa' if completo else 'COM PENDÊNCIAS — lista não confirmada como completa'])
    estado.append(['Documentos validados', len(validos)])
    estado.append(['Documentos pendentes', len(pendentes)])
    for falha in falhas_consulta or []:
        estado.append(['Falha de consulta', falha])
    if not completo:
        wb.active = wb.sheetnames.index('Consulta')
    wb.save(resultado/'instrumentos_consultados_AL.xlsx')
    with (resultado/'instrumentos_consultados_AL.csv').open('w', encoding='utf-8-sig', newline='') as file:
        writer = csv.writer(file, delimiter=';'); writer.writerow(colunas)
        for item in validos+pendentes:
            writer.writerow([("'"+str(item[c]) if str(item.get(c, '')).startswith(('=', '+', '-', '@')) else item.get(c, '')) for c in colunas])
    texto = [f'Registros de {inicio.strftime("%d/%m/%Y")} a {fim.strftime("%d/%m/%Y")} — abrangência em Alagoas',
             'Consulta completa' if completo else 'CONSULTA COM PENDÊNCIAS — não afirmar que a lista está completa',
             f'Documentos lidos e validados: {len(validos)} | pendências de documentos: {len(pendentes)}',
             'Esta lista inclui os registros já existentes no controle.']
    texto.extend(f'{d["Registro"]} | {d["Tipo"]} | registro em {d["Data de registro"]} | {d["Fonte"]}' for d in validos)
    texto.extend(f'PENDENTE: {d.get("Registro", d.get("Solicitação", "sem identificador"))} | {d["Status"]}' for d in pendentes)
    texto.extend('FALHA DE CONSULTA: '+falha for falha in falhas_consulta or [])
    (resultado/'instrumentos_consultados_AL.txt').write_text('\n'.join(texto)+'\n', encoding='utf-8')


def executar_online(controle=None, raiz=Path('saidas'), aba=None, progresso=lambda mensagem: None, ano=None, ontem=False):
    hoje = hoje_alagoas()
    inicio, fim = periodo_consulta(hoje, ano, ontem)
    pasta = raiz / ('online_'+datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:8])
    pasta.mkdir(parents=True, exist_ok=False)
    docs = pasta/'documentos'; docs.mkdir()
    eventos = [f'Consulta iniciada em {datetime.now().isoformat(timespec="seconds")}',
               f'Período de registro: {inicio} a {fim}, fuso America/Maceio; somente CCT/aditivo de CCT; UF AL e todos os municípios de AL em pesquisas separadas']
    instrumentos, pendencias, totais = [], [], {}
    catalogo = []
    try:
        # Valida o controle antes de consultar; entrada é apenas lida.
        monitor.abrir_controle(controle, aba)
        from playwright.sync_api import sync_playwright
        proxy = proxy_do_ambiente()
        with sync_playwright() as p:
            options = opcoes_navegador(proxy)
            eventos.append('Navegador: '+('Microsoft Edge com janela visível; sessão separada' if sys.platform == 'win32'
                                        else 'Chromium sem janela; sessão separada'))
            progresso('Abrindo o navegador para a consulta automática…')
            try:
                browser = p.chromium.launch(**options)
            except Exception as exc:
                if sys.platform == 'win32' and 'distribution' in str(exc) and 'not found' in str(exc):
                    raise ValueError('Microsoft Edge não encontrado. Instale o Edge pelo site oficial da Microsoft e execute novamente. '
                                     'A consulta não foi iniciada') from exc
                raise
            try:
                instrumentos, pendencias, totais = coletar(browser, docs, fim, eventos, progresso, inicio=inicio, catalogo=catalogo)
            finally:
                browser.close()
    except Exception as exc:
        motivo = str(exc) if isinstance(exc, ValueError) else classificar_erro(str(exc))+': '+type(exc).__name__
        pendencias.append(motivo); eventos.append(motivo)
    completo = not pendencias and len(totais) == len(PESQUISAS)
    # Manifesto de auditoria gerado internamente; não é preparado pelo usuário.
    manifesto = pasta/'coleta_automatica.json'
    manifesto.write_text(json.dumps({'consulta_completa': completo, 'instrumentos': instrumentos}, ensure_ascii=False, indent=2), encoding='utf-8')
    resultado = pasta/'resultado'
    code = monitor.executar(manifesto, controle, resultado, fim, aba, modo='online', inicio=inicio)
    eventos.extend('PENDÊNCIA: '+p for p in pendencias)
    log = resultado/'log_convencoes_AL.txt'
    log.write_text('\n'.join(eventos)+'\n'+log.read_text(encoding='utf-8'), encoding='utf-8')
    resumo = resultado/'resumo_convencoes_AL.txt'
    resumo.write_text(resumo.read_text(encoding='utf-8')+'Totais oficiais por pesquisa (podem se sobrepor; não somar): '+str(totais)+'\n'+
                     ('Pendências: '+'; '.join(pendencias)+'\n' if pendencias else ''), encoding='utf-8')
    gerar_relatorio(catalogo, instrumentos, pasta, resultado, inicio, fim, completo and code == 0, pendencias)
    resumo.write_text(resumo.read_text(encoding='utf-8')+f'Instrumentos listados nesta consulta (inclusive já existentes): {len(catalogo)}\nRelatório completo: instrumentos_consultados_AL.xlsx e instrumentos_consultados_AL.txt\n', encoding='utf-8')
    return {'codigo': code, 'pasta': resultado, 'controle': resultado/'controle_convencoes_AL.xlsx',
            'completa': completo and code == 0, 'resumo': resumo.read_text(encoding='utf-8'),
            'relatorio': (resultado/'instrumentos_consultados_AL.txt').read_text(encoding='utf-8')}


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--controle', type=Path)
    parser.add_argument('--aba')
    periodo = parser.add_mutually_exclusive_group()
    periodo.add_argument('--ano', type=int)
    periodo.add_argument('--ontem', action='store_true', help='Somente registros do dia anterior em Alagoas')
    args = parser.parse_args()
    resultado = executar_online(args.controle, aba=args.aba, progresso=print, ano=args.ano, ontem=args.ontem)
    print(resultado['resumo']); print('Resultados:', resultado['pasta'].resolve())
    raise SystemExit(resultado['codigo'])
