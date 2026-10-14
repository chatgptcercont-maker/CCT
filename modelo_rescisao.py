"""Renderiza uma cópia preenchida do modelo, sem os dados do arquivo de exemplo."""
import json
from pathlib import Path
import textwrap
from openpyxl import Workbook
from openpyxl.styles import Font, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject

MODELO=Path(__file__).resolve().parent/'assets/modelo_rescisao.json'


def campos_modelo(recibo):
    from recibo_rescisao import reais
    d=recibo['dados'];v={}
    mapa={'0,5':'documento_empresa','2,5':'empresa','0,7':'endereco_empresa','4,7':'bairro_empresa',
          '0,9':'municipio_empresa','1,9':'uf_empresa','2,9':'cep_empresa','3,9':'cnae','4,9':'tomador',
          '0,12':'pis','2,12':'trabalhador','0,14':'endereco_trabalhador','4,14':'bairro_trabalhador',
          '0,16':'municipio_trabalhador','1,16':'uf_trabalhador','2,16':'cep_trabalhador','3,16':'ctps','5,16':'cpf',
          '0,18':'nascimento','2,18':'mae','0,21':'tipo_contrato','0,23':'causa',
          '0,25':'remuneracao','2,25':'admissao','3,25':'aviso','4,25':'desligamento','5,25':'codigo_afastamento',
          '0,27':'pensao_trct','2,27':'pensao_fgts','4,27':'categoria',
          '0,29':'codigo_sindical','2,29':'entidade_sindical'}
    for coordenada,campo in mapa.items():
        c,r=map(int,coordenada.split(','));v[(r,c)]=d.get(campo,'')
    proventos=[i for i in recibo['itens'] if i['tipo']=='Provento'];descontos=[i for i in recibo['itens'] if i['tipo']=='Desconto']
    if len(proventos)>18 or len(descontos)>12:raise ValueError('O modelo comporta até 18 verbas e 12 deduções. Reduza os itens ou ajuste o modelo.')
    for itens,primeira in [(proventos,34),(descontos,43)]:
        for i,item in enumerate(itens):
            linha=primeira+i//3;coluna=(i%3)*2
            v[(linha,coluna)]=item['descricao'];v[(linha,coluna+1)]=reais(item['valor'])
    v[(40,5)]=reais(recibo['proventos']);v[(47,5)]=reais(recibo['descontos']);v[(48,5)]=reais(recibo['liquido'])
    if d.get('observacoes'):v[(30,0)]='Observações: '+d['observacoes']
    q={(4,0):d.get('documento_empresa',''),(4,2):d['empresa'],(7,0):d.get('pis',''),(7,2):d['trabalhador'],
       (9,0):d.get('ctps',''),(9,2):d.get('cpf',''),(9,3):d.get('nascimento',''),(9,4):d.get('mae',''),
       (12,0):d.get('causa',''),(14,0):d.get('admissao',''),(14,2):d.get('aviso',''),(14,3):d.get('desligamento',''),
       (14,4):d.get('codigo_afastamento',''),(14,5):d.get('pensao_fgts',''),(16,0):d.get('categoria',''),
       (30,0):d['local']+', '+recibo['data'].strftime('%d/%m/%Y')}
    if recibo.get('emitir_quitacao'):
        q[(19,0)]='No dia '+recibo['data_pagamento'].strftime('%d/%m/%Y')+' foi realizado, conforme informado pelo responsável, o efetivo pagamento das'
        q[(20,0)]='verbas rescisórias especificadas no TRCT, no valor líquido de '+reais(recibo['liquido'])+', que é parte integrante deste Termo de Quitação.'
        q[(21,0)]='O termo deve ser conferido e assinado pelas partes.'
    return v,q


def folhas(recibo):
    spec=json.loads(MODELO.read_text(encoding='utf-8'));dados=campos_modelo(recibo)
    resultado=[]
    for i,aba in enumerate(spec['abas']):
        if i==1 and not recibo.get('emitir_quitacao'):continue
        valores={tuple(map(int,k.split(','))):v for k,v in aba['rotulos'].items()};valores.update(dados[i])
        mescladas=[tuple(m) for m in aba['mescladas']]
        extras=[(18,19,0,2),(18,19,2,6)] if i==0 else [(30,31,0,6),(49,50,0,3)]
        for m in extras:
            if not any(m[0]<x[1] and m[1]>x[0] and m[2]<x[3] and m[3]>x[2] for x in mescladas):mescladas.append(m)
        resultado.append((aba,valores,mescladas))
    return resultado


def validar_destino(destino,extensao):
    p=Path(destino)
    if p.suffix.lower()!=extensao:raise ValueError('Escolha um arquivo '+extensao+'.')
    if p.exists():raise ValueError('O arquivo já existe. Escolha outro nome; o original será preservado.')
    return p


def alturas_preenchidas(aba,valores,merges):
    larguras=[x*5.25 for x in aba['larguras']];alturas=list(aba['alturas'])
    for (r,c),valor in valores.items():
        if not str(valor).strip():continue
        m=next((m for m in merges if m[0]<=r<m[1] and m[2]<=c<m[3]),None)
        r2,c2=(m[1],m[3]) if m else (r+1,c+1)
        tamanho=aba['estilos'].get(f'{r},{c}',{}).get('tamanho',10)
        largura=sum(larguras[c:c2])-5
        linhas=textwrap.wrap(str(valor),max(5,int(largura/(tamanho*0.52))),break_long_words=True) or ['']
        necessario=len(linhas)*(tamanho+2)+3
        alturas[r]+=max(0,necessario-sum(alturas[r:r2]))
    return alturas


def exportar_excel(recibo,destino):
    p=validar_destino(destino,'.xlsx');w=Workbook();w.remove(w.active)
    for aba,valores,merges in folhas(recibo):
        s=w.create_sheet(aba['nome'])
        for c,largura in enumerate(aba['larguras'],1):s.column_dimensions[get_column_letter(c)].width=largura
        for r,altura in enumerate(alturas_preenchidas(aba,valores,merges),1):s.row_dimensions[r].height=altura
        for r in range(aba['linhas']):
            for c in range(6):
                st=aba['estilos'].get(f'{r},{c}',{});cel=s.cell(r+1,c+1)
                cel.font=Font(name='Arial',size=st.get('tamanho',10),bold=st.get('negrito',False))
                cel.alignment=Alignment(vertical='center',wrap_text=True)
                if st.get('borda'):cel.border=Border(left=Side(style='thin'),right=Side(style='thin'),top=Side(style='thin'),bottom=Side(style='thin'))
        for (r,c),v in valores.items():
            cel=s.cell(r+1,c+1);cel.value=str(v);cel.data_type='s'  # campos do usuário nunca viram fórmulas
        for r1,r2,c1,c2 in merges:s.merge_cells(start_row=r1+1,end_row=r2,start_column=c1+1,end_column=c2)
        s.page_setup.paperSize=s.PAPERSIZE_A4;s.page_setup.orientation='portrait';s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
        s.sheet_properties.pageSetUpPr.fitToPage=True;s.print_area=f'A1:F{aba["linhas"]}'
    with p.open('xb') as arquivo:w.save(arquivo)
    return p


def exportar_pdf(recibo,destino):
    p=validar_destino(destino,'.pdf');writer=PdfWriter()
    for aba,valores,merges in folhas(recibo):
        larguras=[x*5.25 for x in aba['larguras']];alturas=alturas_preenchidas(aba,valores,merges)
        escala=min(515/sum(larguras),758/sum(alturas));xs=[40];ys=[800]
        for x in larguras:xs.append(xs[-1]+x*escala)
        for y in alturas:ys.append(ys[-1]-y*escala)
        page=writer.add_blank_page(width=595,height=842)
        fontes=DictionaryObject()
        for nome,base in [('F1','Helvetica'),('F2','Helvetica-Bold')]:
            f=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/'+base),NameObject('/Encoding'):NameObject('/WinAnsiEncoding')})
            fontes[NameObject('/'+nome)]=writer._add_object(f)
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):fontes})
        comandos=[b'0.4 w 0 G']
        for r in range(aba['linhas']):
            for c in range(6):
                m=next((m for m in merges if m[0]<=r<m[1] and m[2]<=c<m[3]),None)
                if m and (r,c)!=(m[0],m[2]):continue
                r2,c2=(m[1],m[3]) if m else (r+1,c+1)
                x,y,w,h=xs[c],ys[r2],xs[c2]-xs[c],ys[r]-ys[r2]
                st=aba['estilos'].get(f'{r},{c}',{})
                if st.get('borda'):comandos.append(f'{x:.2f} {y:.2f} {w:.2f} {h:.2f} re S'.encode())
                valor=str(valores.get((r,c),'')).strip()
                if not valor:continue
                tamanho=st.get('tamanho',10)*escala;fonte='F2' if st.get('negrito') else 'F1'
                if tamanho<4.5:raise ValueError('Conteúdo extenso demais para uma página legível. Abrevie os campos do modelo.')
                linhas=textwrap.wrap(valor,max(5,int((w-5)/(tamanho*0.52))),break_long_words=True) or ['']
                while len(linhas)*(tamanho+1)>h-1 and tamanho>4.5:
                    tamanho-=0.25
                    linhas=textwrap.wrap(valor,max(5,int((w-5)/(tamanho*0.52))),break_long_words=True) or ['']
                if len(linhas)*(tamanho+1)>h-1:
                    raise ValueError('Texto longo demais para o campo do modelo: '+valor[:35]+'. Abrevie o campo; nenhuma informação foi cortada.')
                for n,linha in enumerate(linhas):
                    try:raw=linha.encode('cp1252')
                    except UnicodeEncodeError as exc:raise ValueError('Texto com caracteres não suportados pelo PDF.') from exc
                    raw=raw.replace(b'\\',b'\\\\').replace(b'(',b'\\(').replace(b')',b'\\)')
                    comando=f'BT /{fonte} {tamanho:.2f} Tf {x+2:.2f} {ys[r]-tamanho-1-n*(tamanho+1):.2f} Td '.encode()
                    comandos.append(comando+b'('+raw+b') Tj ET')
        stream=DecodedStreamObject();stream.set_data(b'\n'.join(comandos));page[NameObject('/Contents')]=writer._add_object(stream)
    with p.open('xb') as arquivo:writer.write(arquivo)
    return p
