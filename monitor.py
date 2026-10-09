"""Monitor de convenções: consulta online ou alternativa assistida, preservando o controle original."""
import argparse
from datetime import date, datetime, timedelta
import hashlib
import json
from pathlib import Path
import re
import unicodedata
import urllib.error
import urllib.request
import uuid

from openpyxl import Workbook, load_workbook
from pypdf import PdfReader

CONSULTA = 'https://mediador.trabalho.gov.br/sistemas/mediador/ConsultarInstColetivo'
COLUNAS = ['Registro', 'Data de registro', 'Tipo', 'Abrangência territorial', 'Documento', 'Fonte', 'SHA256']
REGEX_REGISTRO = r'(?:SRT|AC|AL|AM|AP|BA|CE|DF|ES|GO|MA|MG|MS|MT|PA|PB|PE|PI|PR|RJ|RN|RO|RR|RS|SC|SE|SP|TO)\d{6}/\d{4}'


def normalizar(texto):
    return ''.join(c for c in unicodedata.normalize('NFD', str(texto)) if unicodedata.category(c) != 'Mn').upper()


def registro(valor):
    valor = re.sub(r'\s+', '', normalizar(valor))
    if not re.fullmatch(REGEX_REGISTRO, valor):
        raise ValueError('Número de registro ausente ou inválido')
    return valor


def extrair(texto, hoje):
    """Campos vêm exclusivamente do texto do documento; dúvida impede inclusão."""
    n = normalizar(texto)
    titulo = n[:3000]
    if re.search(r'(?:TERMO ADITIVO (?:DE|A|AO) )?ACORDO COLETIVO', titulo):
        raise ValueError('Excluído: acordo coletivo ou aditivo de acordo')
    if re.search(r'TERMO ADITIVO (?:DE|A|A UMA|A CONVENCAO)?\s*CONVENCAO COLETIVA', titulo):
        tipo = 'Termo Aditivo de Convenção Coletiva'
    elif 'TERMO ADITIVO' in titulo:
        raise ValueError('Tipo de aditivo não confirmado')
    elif 'CONVENCAO COLETIVA DE TRABALHO' in titulo:
        tipo = 'Convenção Coletiva de Trabalho'
    else:
        raise ValueError('Tipo de instrumento não confirmado')
    regs = re.findall(r'NUMERO DE REGISTRO NO MTE\s*[:\-]?\s*(' + REGEX_REGISTRO + r')', n)
    datas = re.findall(r'DATA DE REGISTRO NO MTE\s*[:\-]?\s*(\d{2}/\d{2}/\d{4})', n)
    if len(set(regs)) != 1 or len(set(datas)) != 1:
        raise ValueError('Registro ou data ambíguos/ausentes; conferir documento')
    data = datetime.strptime(datas[0], '%d/%m/%Y').date()
    if not hoje - timedelta(days=6) <= data <= hoje:
        raise ValueError('Excluído: data de registro fora dos sete dias inclusivos')
    # A UF do número de registro não prova abrangência territorial.
    trecho = re.search(r'(?:^|\n)\s*(?:CLAUSULA[^\n]*?[-–]\s*)?ABRANGENCIA(?: TERRITORIAL)?[ \t:]*\n(.*?)(?=\n\s*(?:CLAUSULA|SALARIOS|REMUNERACAO)|\Z)', n, re.S)
    if not trecho or not re.search(r'\bALAGOAS\b|(?:/|\bUF\s*[:=]?)\s*AL\b|\bAL\b', trecho[1]):
        raise ValueError('Abrangência em Alagoas não confirmada na cláusula territorial')
    return {'Registro': registro(regs[0]), 'Data de registro': data.strftime('%d/%m/%Y'),
            'Tipo': tipo, 'Abrangência territorial': re.sub(r'\s+', ' ', trecho[1]).strip()}


def ler_pdf(caminho):
    if caminho.suffix.lower() != '.pdf':
        raise ValueError('Documento deve ser PDF original')
    if caminho.stat().st_size > 30 * 1024 * 1024:
        raise ValueError('PDF excede 30 MB; conferir manualmente')
    pdf = PdfReader(caminho)
    if pdf.is_encrypted:
        raise ValueError('PDF protegido; leitura indisponível')
    texto = '\n'.join(p.extract_text() or '' for p in pdf.pages)
    if len(texto.strip()) < 80:
        raise ValueError('PDF sem texto legível; pode precisar de OCR e conferência manual')
    return texto


def fonte_oficial(url):
    from urllib.parse import urlsplit
    u = urlsplit(url)
    return u.scheme == 'https' and u.hostname in {'mediador.trabalho.gov.br', 'www3.mte.gov.br', 'www.mte.gov.br', 'mediador.mte.gov.br'} and not u.username and not u.password


def abrir_controle(caminho, aba=None):
    if caminho is None:
        wb = Workbook()
        ws = wb.active
        ws.title = 'Convenções AL'
        ws.append(COLUNAS)
    else:
        if caminho.suffix.lower() != '.xlsx':
            raise ValueError('Controle deve ser .xlsx')
        wb = load_workbook(caminho, data_only=False)
        if aba:
            if aba not in wb.sheetnames:
                raise ValueError('Aba informada não existe')
            ws = wb[aba]
        elif len(wb.sheetnames) == 1:
            ws = wb.active
        else:
            raise ValueError('Controle com várias abas: informe --aba')
    headers = [c.value for c in ws[1]]
    if any(headers.count(c) != 1 for c in COLUNAS):
        raise ValueError('Cabeçalhos incompatíveis. Use o modelo documentado; não adaptar a planilha oficial automaticamente')
    idx = headers.index('Registro')
    existentes = set()
    for row in ws.iter_rows(min_row=2):
        value = row[idx].value
        if value is not None and str(value).strip():
            existentes.add(registro(value))
    return wb, ws, headers, existentes


def executar(manifesto, controle, saida, hoje, aba=None, modo="assistido"):
    saida.mkdir(parents=True, exist_ok=False)
    eventos, novos = [], []
    erro = False
    try:
        wb, ws, headers, existentes = abrir_controle(controle, aba)
        payload = json.loads(manifesto.read_text(encoding='utf-8'))
        if not isinstance(payload, dict) or not isinstance(payload.get('instrumentos'), list):
            raise ValueError('Manifesto deve conter a lista instrumentos')
        completo = payload.get('consulta_completa') is True
        eventos.append(('Coleta online concluída: ' if modo == 'online' else 'Consulta manual declarada completa: ') + str(completo) + ('; verificar pendências de coleta' if modo == 'online' else '; declaração não verificada pelo programa'))
        if not completo:
            erro = True
        for i, item in enumerate(payload['instrumentos'], 1):
            try:
                if not isinstance(item, dict):
                    raise ValueError('Item inválido')
                if item.get('status', 'disponivel') != 'disponivel':
                    raise ValueError('Pendência informada: ' + str(item.get('status')))
                fonte = item.get('fonte', '')
                if not fonte_oficial(fonte):
                    raise ValueError('Informe o endereço HTTPS original do Mediador/MTE')
                relativo = Path(item.get('arquivo', ''))
                base = manifesto.parent.resolve()
                caminho = (base / relativo).resolve()
                if relativo.is_absolute() or not caminho.is_relative_to(base):
                    raise ValueError('PDF deve estar dentro da pasta do manifesto')
                dados = extrair(ler_pdf(caminho), hoje)
                if item.get('registro') and registro(item['registro']) != dados['Registro']:
                    raise ValueError('Registro do resultado diverge do PDF')
                if dados['Registro'] in existentes:
                    eventos.append(f'Item {i}: já existente {dados["Registro"]}')
                    continue
                dados.update({'Documento': relativo.as_posix(), 'Fonte': fonte,
                              'SHA256': hashlib.sha256(caminho.read_bytes()).hexdigest()})
                # Evita interpretação de conteúdo de PDF como fórmula de Excel.
                ws.append([("'" + dados[c] if dados.get(c, '').startswith(('=', '+', '-', '@')) else dados.get(c)) for c in headers])
                existentes.add(dados['Registro'])
                novos.append(dados)
                eventos.append(f'Item {i}: incluído {dados["Registro"]}')
            except Exception as exc:
                mensagem = str(exc)
                if not mensagem.startswith('Excluído:'):
                    erro = True
                eventos.append(f'Item {i}: {type(exc).__name__}: {mensagem}')
        wb.save(saida / 'controle_convencoes_AL.xlsx')
    except Exception as exc:
        erro = True
        novos.clear()
        eventos.append(f'Falha geral: {type(exc).__name__}: {exc}; planilha de saída não confirmada')
    resumo = [f'Período de registro: {(hoje-timedelta(days=6)).strftime("%d/%m/%Y")} a {hoje.strftime("%d/%m/%Y")} (sete dias inclusivos)',
              ('Modo online: documentos obtidos automaticamente na fonte oficial; verificar log e conferir a cópia do controle.' if modo == 'online' else 'Modo assistido: PDFs fornecidos pelo usuário; autenticidade e completude da busca exigem conferência humana.'),
              'Resultado: ' + ('COM PENDÊNCIAS' if erro else 'PROCESSAMENTO CONCLUÍDO'), f'Novos instrumentos: {len(novos)}']
    resumo.extend(f'{d["Registro"]} | {d["Tipo"]} | registro em {d["Data de registro"]} | {d["Fonte"]}' for d in novos)
    resumo.append('Ausência de inclusões não comprova ausência de instrumentos no Mediador.')
    (saida / 'resumo_convencoes_AL.txt').write_text('\n'.join(resumo)+'\n', encoding='utf-8')
    (saida / 'log_convencoes_AL.txt').write_text('\n'.join(eventos)+'\n', encoding='utf-8')
    return 2 if erro else 0


def consultar(saida):
    """Diagnóstico limitado de acesso público, sem submissão de filtros ou CAPTCHA."""
    saida.mkdir(parents=True, exist_ok=False)
    try:
        req = urllib.request.Request(CONSULTA, headers={'User-Agent': 'CercontMonitor/0.1 (diagnostico manual)'})
        with urllib.request.urlopen(req, timeout=25) as r:
            texto = normalizar(r.read(2_000_000).decode('utf-8', errors='replace'))
            evento = f'HTTP {r.status}. '
            if any(s in texto for s in ['CAPTCHA', 'RECAPTCHA', 'HCAPTCHA']):
                evento += 'CAPTCHA detectado; use consulta manual, sem contornar proteção.'
            else:
                evento += 'Página recebida; consulta automatizada ainda não validada. Use modo assistido.'
    except (urllib.error.URLError, TimeoutError, OSError) as exc:
        evento = f'Acesso bloqueado ou indisponível: {type(exc).__name__}: {exc}'
    (saida/'log_convencoes_AL.txt').write_text(evento+'\n', encoding='utf-8')
    (saida/'resumo_convencoes_AL.txt').write_text('CONSULTA NÃO CONCLUÍDA. Nenhum resultado inferido.\n'+evento+'\n', encoding='utf-8')
    return 2


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--manifesto', type=Path)
    p.add_argument('--controle', type=Path, help='Planilha de entrada, somente leitura')
    p.add_argument('--aba')
    p.add_argument('--saida', type=Path, help='Assistido: pasta nova; online: raiz das pastas de execução')
    p.add_argument('--hoje', type=date.fromisoformat, help='Somente modo assistido: AAAA-MM-DD; padrão: data local')
    p.add_argument('--consultar', action='store_true', help='Consulta online automática com navegador, sem PDFs ou manifesto do usuário')
    args = p.parse_args()
    if not args.consultar and not args.manifesto:
        p.error('Informe --manifesto ou --consultar')
    if args.consultar and args.manifesto:
        p.error('--consultar não combina com --manifesto')
    if args.consultar and args.hoje is not None:
        p.error('--hoje é exclusivo do modo assistido; a consulta online usa a data atual de Alagoas')
    if args.consultar:
        from consulta_online import executar_online
        resultado = executar_online(args.controle, args.saida or Path('saidas'), args.aba, print)
        print(resultado['resumo'])
        print('Resultados:', resultado['pasta'].resolve())
        return resultado['codigo']
    saida = args.saida or Path('saidas') / (datetime.now().strftime('%Y%m%d_%H%M%S')+'_'+uuid.uuid4().hex[:8])
    try:
        code = executar(args.manifesto, args.controle, saida, args.hoje or date.today(), args.aba)
    except FileExistsError:
        print('Pasta de saída já existe. Escolha uma pasta nova; nenhum arquivo foi sobrescrito.')
        return 2
    print(f'Resultados: {saida.resolve()} | código {code}')
    return code


if __name__ == '__main__':
    raise SystemExit(main())
